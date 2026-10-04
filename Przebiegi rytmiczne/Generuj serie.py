import io
import os
import re
import shutil
import sys
import zipfile
from collections import Counter
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
import yaml

FOLDER_PRZEBIEGOW_INPUT = "Definicje serii"
FOLDER_ODPOWIEDZI = "Odpowiedzi - tylko dla nauczyciela"
FOLDER_DLA_UCZNIOW = "Dla uczniów - do skopiowania na chmurę"
FOLDER_WERSJI = "Wersje"
SZABLON_PRZYDZIALOW = "Przydzialy/przydzialy_klasa_{klasa}.txt"

KATALOG_SKRYPTU = os.path.dirname(os.path.abspath(__file__))


class StrumienKonsoliZLogiem:
    def __init__(self, strumien_ekranu, bufor_logu):
        self.strumien_ekranu = strumien_ekranu
        self.bufor_logu = bufor_logu
        self.plik_logu = None

    def write(self, tekst):
        self.strumien_ekranu.write(tekst)
        self.strumien_ekranu.flush()
        if self.plik_logu:
            self.plik_logu.write(tekst)
            self.plik_logu.flush()
        else:
            self.bufor_logu.write(tekst)
        return len(tekst)

    def flush(self):
        self.strumien_ekranu.flush()
        if self.plik_logu:
            self.plik_logu.flush()

    def podlacz_plik_logu(self, plik_logu):
        self.plik_logu = plik_logu

    def isatty(self):
        return self.strumien_ekranu.isatty()


def wypisz_blad_na_czerwono(error):
    print(
        f"\033[91mNie udało się wygenerować materiałów. "
        f"Szczegóły: {error}\033[0m",
        file=sys.stderr,
    )


def wczytaj_konfiguracje(nazwa_pliku=".przebiegi rytmiczne.yml"):
    sciezka_pliku = os.path.join(KATALOG_SKRYPTU, nazwa_pliku)

    if not os.path.exists(sciezka_pliku):
        raise FileNotFoundError(
            f"Nie znaleziono pliku konfiguracyjnego pod ścieżką: {sciezka_pliku}"
        )

    with open(sciezka_pliku, "r", encoding="utf-8") as file:
        return yaml.safe_load(file)


def wczytaj_uczniow_z_pliku(sciezka_pliku):
    with open(sciezka_pliku, "r", encoding="utf-8") as f:
        uczniowie = [line.strip() for line in f if line.strip()]
    return uczniowie


def generuj_zestaw_rownomierny(liczba_uczniow, numery_przebiegow):
    """
    Deterministycznie i równomiernie przydziela podane numery przebiegów.
    Gwarantuje, że różnica w liczbie wystąpień między jakimikolwiek dwoma numerami
    nie przekroczy 1.
    """
    zakres = list(numery_przebiegow)
    if not zakres:
        raise ValueError("Nie można wygenerować przydziałów bez numerów przebiegów.")
    len_z = len(zakres)

    wymagane_numery = liczba_uczniow * 2
    sekwencja_numerow = [zakres[i % len_z] for i in range(wymagane_numery)]

    zestawy = []
    uzyte_pary = set()

    for i in range(liczba_uczniow):
        l1 = sekwencja_numerow[2 * i]
        l2 = sekwencja_numerow[2 * i + 1]

        if l1 == l2 or tuple(sorted((l1, l2))) in uzyte_pary:
            offset = 1
            while (l1 == l2 or tuple(sorted((l1, l2))) in uzyte_pary) and (
                2 * i + 1 + offset
            ) < len(sekwencja_numerow):
                idx_alt = 2 * i + 1 + offset
                sekwencja_numerow[2 * i + 1], sekwencja_numerow[idx_alt] = (
                    sekwencja_numerow[idx_alt],
                    sekwencja_numerow[2 * i + 1],
                )
                l2 = sekwencja_numerow[2 * i + 1]
                offset += 1

        para = tuple(sorted((l1, l2)))
        uzyte_pary.add(para)
        zestawy.append([l1, l2])

    return zestawy


def znajdz_plik_klasy(folder_klas, nazwa_klasy):
    """
    Szuka pliku z listą uczniów w folderze klas.
    Sprawdza kolejno: nazwa.txt, Klasa_nazwa.txt, klasa_nazwa.txt
    """
    mozliwe_nazwy = [
        f"{nazwa_klasy}.txt",
        f"Klasa_{nazwa_klasy}.txt",
        f"klasa_{nazwa_klasy}.txt",
    ]

    for nazwa_pliku in mozliwe_nazwy:
        sciezka = os.path.join(folder_klas, nazwa_pliku)
        if os.path.exists(sciezka):
            return sciezka

    return None


def przydziel_numery_dla_klas(konfiguracja, folder_wynikowy, folder_wersji):
    numery_przebiegow = konfiguracja["numery_przebiegow"]
    liczba_przebiegow = len(numery_przebiegow)

    rel_folder_klas = konfiguracja.get("folder_klas", "../Listy uczniow")
    folder_klas = os.path.abspath(
        os.path.join(KATALOG_SKRYPTU, rel_folder_klas)
    )

    klasy = konfiguracja.get("klasy", [])
    if not klasy:
        print("Brak zdefiniowanych klas w pliku konfiguracyjnym (sekcja 'klasy').")
        return []

    polowa = liczba_przebiegow // 2
    pliki_przydzialow = []

    for nazwa_klasy in klasy:
        nazwa_klasy_str = str(nazwa_klasy).strip()
        
        # Oczyszczenie nazwy klasy z przedrostka Klasa_ na potrzeby zmiennej {klasa}
        czysta_nazwa_klasy = nazwa_klasy_str.replace("Klasa_", "").replace("klasa_", "")

        sciezka_pliku_wejsciowego = znajdz_plik_klasy(folder_klas, czysta_nazwa_klasy)

        if not sciezka_pliku_wejsciowego:
            print(
                f"Ostrzeżenie: Nie znaleziono pliku z listą uczniów dla klasy '{czysta_nazwa_klasy}' w folderze: {folder_klas}"
            )
            continue

        uczniowie = wczytaj_uczniow_z_pliku(sciezka_pliku_wejsciowego)

        if not uczniowie:
            print(f"Plik dla klasy '{czysta_nazwa_klasy}' jest pusty. Pomijam.")
            continue

        n = len(uczniowie)

        zestawy_A = generuj_zestaw_rownomierny(
            n, numery_przebiegow[:polowa]
        )
        zestawy_B = generuj_zestaw_rownomierny(
            n, numery_przebiegow[polowa:]
        )

        zawartosc_pliku = []

        for idx, uczeń in enumerate(uczniowie, start=1):
            A = zestawy_A[idx - 1]
            B = zestawy_B[idx - 1]

            if idx % 2 != 0:
                wynikowe = [A[0], B[0], A[1], B[1]]
            else:
                wynikowe = [B[0], A[0], B[1], A[1]]

            sformatowane_liczby = [f"{num:02d}" for num in wynikowe]
            ciag_liczb = ", ".join(sformatowane_liczby)

            uczen_format = f"{uczeń}:".ljust(35)
            linia = f"{uczen_format}{ciag_liczb}"
            zawartosc_pliku.append(linia)

        # Tworzenie ścieżki docelowej z podstawieniem czystej nazwy klasy (np. "5")
        wzgledna_sciezka_wynikowa = SZABLON_PRZYDZIALOW.replace("\\", "/").format(
            klasa=czysta_nazwa_klasy
        )
        sciezka_pliku_wynikowego = os.path.abspath(
            os.path.join(folder_wersji, wzgledna_sciezka_wynikowa)
        )

        os.makedirs(os.path.dirname(sciezka_pliku_wynikowego), exist_ok=True)

        with open(sciezka_pliku_wynikowego, "w", encoding="utf-8") as f:
            f.write("\n".join(zawartosc_pliku))
        pliki_przydzialow.append(
            (sciezka_pliku_wynikowego, f"Klasa_{czysta_nazwa_klasy}")
        )

        print(
            f"\nPrzetworzono klasę '{czysta_nazwa_klasy}' ({n} uczniów) -> Zapisano w '{sciezka_pliku_wynikowego}'"
        )

    return pliki_przydzialow


def wypisz_raport_koncowy(
    numery_przebiegow, pliki_klas, folder_dla_uczniow
):
    """Wyświetla wyłącznie statystyki przebiegów i przydziałów klas."""
    raport = [
        "STATYSTYKI KOŃCOWE",
        f"Liczba przebiegów: {len(numery_przebiegow)}",
        f"Pliki do skopiowania dla uczniów: {folder_dla_uczniow}",
    ]

    for sciezka_txt, nazwa_klasy in pliki_klas:
        if len(raport) > 2:
            raport.append("")
        with open(sciezka_txt, "r", encoding="utf-8") as file:
            linie = [line.strip() for line in file if line.strip()]

        liczba_uczniow = 0
        licznik_numerow = Counter()
        liczniki_pozycji = [Counter() for _ in range(4)]
        for linia in linie:
            if ":" not in linia:
                continue
            _, numery_str = linia.split(":", 1)
            numery = [
                int(value.strip())
                for value in numery_str.split(",")
                if value.strip().isdigit()
            ]
            liczba_uczniow += 1
            licznik_numerow.update(numery)
            for pozycja, numer in enumerate(numery[:4]):
                liczniki_pozycji[pozycja][numer] += 1

        raport.append(nazwa_klasy)
        raport.append(f"  Uczniów: {liczba_uczniow}")
        licznik_czestotliwosci = Counter(
            licznik_numerow.get(numer, 0) for numer in numery_przebiegow
        )
        grupy_rozkładu = []
        for ile_razy, liczba_numerow in sorted(
            licznik_czestotliwosci.items()
        ):
            grupy_rozkładu.append(
                f"    {liczba_numerow} przebiegów {ile_razy} razy powtórzonych"
            )
        raport.append("  Częstotliwość:")
        raport.extend(grupy_rozkładu)
        raport.append(
            "  Numery występujące u więcej niż jednego ucznia "
            "na tej samej pozycji:"
        )
        for pozycja, licznik_pozycji in enumerate(liczniki_pozycji, start=1):
            powtarzajace_sie_numery = sorted(
                numer
                for numer, liczba_wystapien in licznik_pozycji.items()
                if liczba_wystapien > 1
            )
            lista_numerow = (
                ", ".join(
                    f"{numer:02d}" for numer in powtarzajace_sie_numery
                )
                if powtarzajace_sie_numery
                else "brak"
            )
            raport.append(
                f"    Pozycja {pozycja}: "
                f"{len(powtarzajace_sie_numery)} powtarzających się numerów "
                f"({lista_numerow})"
            )

    print(f"\033[92m{'\n'.join(raport)}\033[0m")


def znajdz_plik_word(path_input):
    """Zwraca jedyny plik DOCX w serii i zgłasza błąd dla kilku plików."""
    pliki_docx = []
    for root, dirs, files in os.walk(path_input):
        for nazwa_pliku in files:
            if nazwa_pliku.lower().endswith(".docx") and not nazwa_pliku.startswith("~$"):
                pliki_docx.append(os.path.join(root, nazwa_pliku))

    if len(pliki_docx) >= 2:
        lista_plikow = "\n".join(pliki_docx)
        raise ValueError(
            f"Błąd: znaleziono {len(pliki_docx)} pliki Word (.docx) w serii "
            f"'{path_input}'. Maksymalna liczba to 1:\n{lista_plikow}"
        )

    if not pliki_docx:
        raise FileNotFoundError(
            f"Nie znaleziono pliku Word (.docx) w serii: {path_input}"
        )

    return pliki_docx[0]


def pobierz_przebiegi_word(sciezka_docx):
    """Odczytuje numery przebiegów z pierwszej kolumny tabel i obrazy z drugiej."""
    doc = Document(sciezka_docx)
    rid_to_path = {
        rel.rId: rel.target_ref
        for rel in doc.part.rels.values()
        if "image" in rel.target_ref
    }
    przebiegi = []
    uzyte_numery = set()

    for numer_tabeli, table in enumerate(doc.tables, start=1):
        for numer_wiersza, row in enumerate(table.rows, start=1):
            if len(row.cells) < 2:
                continue
            image_rids = row.cells[1]._tc.xpath(".//a:blip/@r:embed")
            dopasowanie = re.search(r"\d+", row.cells[0].text)
            if not dopasowanie:
                if image_rids:
                    raise ValueError(
                        f"W tabeli {numer_tabeli}, wierszu {numer_wiersza} "
                        "znaleziono zapis nutowy bez numeru przebiegu "
                        "w pierwszej kolumnie."
                    )
                continue

            numer = int(dopasowanie.group())
            if numer in uzyte_numery:
                raise ValueError(
                    f"Powtórzony numer przebiegu w dokumencie Word: {numer}."
                )

            if not image_rids:
                raise ValueError(
                    f"Brak obrazu w drugim polu dla przebiegu {numer}."
                )

            target = rid_to_path.get(image_rids[0])
            if not target:
                raise ValueError(
                    f"Nie można odczytać obrazu dla przebiegu {numer}."
                )
            if not target.startswith("word/"):
                target = "word/" + target

            przebiegi.append((numer, target))
            uzyte_numery.add(numer)

    if not przebiegi:
        raise ValueError(
            f"Nie znaleziono numerowanych przebiegów w tabelach dokumentu: {sciezka_docx}"
        )

    return przebiegi


def wyodrebnij_obrazy_przebiegow(sciezka_docx, path_przebiegi_img, przebiegi):
    """Wypakowuje przekazane obrazy przebiegów z dokumentu Word."""
    os.makedirs(path_przebiegi_img, exist_ok=True)
    print(f"Przetwarzanie pliku z obrazami przebiegów: {sciezka_docx}")
    with zipfile.ZipFile(sciezka_docx, "r") as archive:
        for numer, img_path in przebiegi:
            ext = os.path.splitext(img_path)[1]
            out_filename = os.path.join(
                path_przebiegi_img, f"{numer:02d}{ext}"
            )
            with open(out_filename, "wb") as output_file:
                output_file.write(archive.read(img_path))
            print(
                f"  [Nutowe] Przebieg {numer:02d} -> wyciągnięto obraz: "
                f"{os.path.basename(img_path)}"
            )

    print(f"Wyodrębniono {len(przebiegi)} obrazów przebiegów.")
    return len(przebiegi)


def znajdz_pliki_audio(path_input, numery_przebiegow):
    """Wyszukuje nagrania i sprawdza kompletność przed rozpoczęciem generowania."""
    audio_map = {}
    numery_do_wyszukania = set(numery_przebiegow)

    def rozpoznaj_numer(nazwa_pliku):
        nazwa_bez_ext = os.path.splitext(nazwa_pliku)[0]
        if nazwa_bez_ext.strip().isdigit():
            numer = int(nazwa_bez_ext.strip())
            if numer in numery_do_wyszukania:
                return numer

        dopasowanie = re.search(
            r"(?:przebieg|nagranie|nr|zestaw)[\s_|-]*0*(\d{1,2})",
            nazwa_bez_ext,
            re.IGNORECASE,
        )
        if dopasowanie:
            numer = int(dopasowanie.group(1))
            if numer in numery_do_wyszukania:
                return numer

        for numer_str in re.findall(r"\d+", nazwa_bez_ext):
            numer = int(numer_str)
            if numer in numery_do_wyszukania:
                return numer
        return None

    print("\n--- Sprawdzanie kompletności nagrań audio MP3 ---")
    for root, dirs, files in os.walk(path_input):
        for nazwa_pliku in files:
            ext = os.path.splitext(nazwa_pliku)[1].lower()
            sciezka = os.path.join(root, nazwa_pliku)
            if ext == ".mp3":
                numer = rozpoznaj_numer(nazwa_pliku)
                if numer is not None and numer not in audio_map:
                    audio_map[numer] = ("plik", sciezka, nazwa_pliku)
            elif ext == ".zip":
                try:
                    with zipfile.ZipFile(sciezka, "r") as archive:
                        for member in archive.namelist():
                            if not member.lower().endswith(".mp3"):
                                continue
                            nazwa_audio = os.path.basename(member)
                            numer = rozpoznaj_numer(nazwa_audio)
                            if numer is not None and numer not in audio_map:
                                audio_map[numer] = (
                                    "zip",
                                    sciezka,
                                    member,
                                    nazwa_audio,
                                )
                except Exception as error:
                    print(f"Błąd podczas odczytu ZIP {nazwa_pliku}: {error}")

    brakujace_numery = [
        numer for numer in numery_przebiegow if numer not in audio_map
    ]
    if brakujace_numery:
        lista_brakujacych = ", ".join(
            f"{numer:02d}" for numer in brakujace_numery
        )
        raise FileNotFoundError(
            f"Brak nagrań MP3 dla przebiegów: {lista_brakujacych}. "
            "Przerywam bez generowania wyników."
        )

    print(
        f"Znaleziono komplet nagrań: "
        f"{len(audio_map)}/{len(numery_przebiegow)}.\n"
    )
    return audio_map


def przygotuj_pliki_audio(audio_map, path_przebiegi_audio):
    """Kopiuje wcześniej zweryfikowane nagrania do folderu wersji."""
    os.makedirs(path_przebiegi_audio, exist_ok=True)
    for numer, zrodlo in sorted(audio_map.items()):
        sciezka_docelowa = os.path.join(
            path_przebiegi_audio, f"{numer:02d}.mp3"
        )
        if zrodlo[0] == "plik":
            shutil.copy(zrodlo[1], sciezka_docelowa)
            nazwa_audio = zrodlo[2]
            typ_zrodla = "Audio"
        else:
            with zipfile.ZipFile(zrodlo[1], "r") as archive:
                with open(sciezka_docelowa, "wb") as output_file:
                    output_file.write(archive.read(zrodlo[2]))
            nazwa_audio = zrodlo[3]
            typ_zrodla = "Audio ZIP"
        print(
            f"  [{typ_zrodla}] Przebieg {numer:02d} -> '{nazwa_audio}'"
        )
    print(f"Przygotowano {len(audio_map)} nagrań audio.\n")


def przetworz_klasy(
    pliki_klas, path_output, path_przebiegi_img, path_przebiegi_audio
):
    """Tworzy dokumenty klas oraz foldery uczniów z przypisanymi nagraniami."""
    if not pliki_klas:
        print("Brak wygenerowanych plików przydziałów do przetworzenia.")
        return

    nazwa_serii = os.path.basename(os.path.normpath(path_output))
    obrazy_przebiegow = {}
    for nazwa_pliku in os.listdir(path_przebiegi_img):
        nazwa, ext = os.path.splitext(nazwa_pliku)
        if nazwa.isdigit():
            obrazy_przebiegow[int(nazwa)] = os.path.join(
                path_przebiegi_img, nazwa_pliku
            )

    nagrania_audio = {}
    if os.path.exists(path_przebiegi_audio):
        for nazwa_pliku in os.listdir(path_przebiegi_audio):
            nazwa, ext = os.path.splitext(nazwa_pliku)
            if nazwa.isdigit() and ext.lower() == ".mp3":
                nagrania_audio[int(nazwa)] = os.path.join(
                    path_przebiegi_audio, nazwa_pliku
                )

    for sciezka_txt, nazwa_klasy in pliki_klas:
        folder_klasy_out = os.path.join(
            path_output, FOLDER_DLA_UCZNIOW, nazwa_klasy
        )
        os.makedirs(folder_klasy_out, exist_ok=True)
        folder_odpowiedzi = os.path.join(path_output, FOLDER_ODPOWIEDZI)
        os.makedirs(folder_odpowiedzi, exist_ok=True)
        doc = Document()
        print(f"\n--- Przetwarzanie: {nazwa_klasy} ---")

        with open(sciezka_txt, "r", encoding="utf-8") as file:
            linie = [line.strip() for line in file if line.strip()]

        for idx, linia in enumerate(linie):
            if ":" not in linia:
                continue
            uczen_info, numery_str = linia.split(":", 1)
            nazwa_ucznia = uczen_info.strip()
            nazwa_ucznia_bez_nr = re.sub(r"^\d+\.\s*", "", nazwa_ucznia).strip()
            numery = [
                int(value.strip())
                for value in numery_str.split(",")
                if value.strip().isdigit()
            ]

            folder_ucznia = os.path.join(folder_klasy_out, nazwa_ucznia)
            podfolder_audio = os.path.join(
                folder_ucznia, nazwa_ucznia_bez_nr, nazwa_serii
            )
            os.makedirs(podfolder_audio, exist_ok=True)
            kopiowane_audio_count = 0
            for numer_porzadkowy, numer_przebiegu in enumerate(numery, start=1):
                if numer_przebiegu in nagrania_audio:
                    shutil.copy(
                        nagrania_audio[numer_przebiegu],
                        os.path.join(podfolder_audio, f"{numer_porzadkowy}.mp3"),
                    )
                    kopiowane_audio_count += 1
            print(
                f" Uczeń: {nazwa_ucznia} -> przekopiowano "
                f"{kopiowane_audio_count}/{len(numery)} plików audio."
            )
            if len(numery) != 4 or kopiowane_audio_count != 4:
                raise ValueError(
                    f"Uczeń '{nazwa_ucznia}' ({nazwa_klasy}): "
                    f"skopiowano {kopiowane_audio_count}/4 plików audio "
                    f"przy {len(numery)} przypisanych numerach."
                )

            doc.add_heading(nazwa_ucznia, level=2)
            table = doc.add_table(rows=0, cols=2)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            for numer_przebiegu in numery:
                row_cells = table.add_row().cells
                row_cells[0].width = Inches(0.8)
                row_cells[1].width = Inches(5.2)
                paragraph_number = row_cells[0].paragraphs[0]
                paragraph_number.alignment = WD_ALIGN_PARAGRAPH.CENTER
                number_run = paragraph_number.add_run(f"{numer_przebiegu:02d}.")
                number_run.font.bold = True
                number_run.font.size = Pt(14)
                paragraph_content = row_cells[1].paragraphs[0]
                if numer_przebiegu in obrazy_przebiegow:
                    paragraph_content.add_run().add_picture(
                        obrazy_przebiegow[numer_przebiegu], width=Inches(5.0)
                    )
                else:
                    paragraph_content.add_run(
                        f"[Brak obrazka dla nr {numer_przebiegu:02d}]"
                    )

            if idx < len(linie) - 1:
                doc.add_page_break()

        sciezka_doc_out = os.path.join(
            folder_odpowiedzi, f"{nazwa_klasy}.docx"
        )
        doc.save(sciezka_doc_out)
        print(f"Pomyślnie wygenerowano komplet dla klasy: {nazwa_klasy}")


def wybierz_serie(path_input):
    """Pyta o serię i zwraca ścieżkę do wybranego podfolderu."""
    serie = [
        nazwa
        for nazwa in os.listdir(path_input)
        if os.path.isdir(os.path.join(path_input, nazwa))
    ]
    serie.sort(
        key=lambda nazwa: (0, int(nazwa))
        if nazwa.isdecimal()
        else (1, nazwa.casefold())
    )

    if not serie:
        raise FileNotFoundError(
            f"Nie znaleziono folderów serii w katalogu: {path_input}"
        )

    print(f"Dostępne serie: {', '.join(serie)}")
    while True:
        wybor = input("Którą serię przetworzyć? Podaj jej nazwę: ").strip()
        if wybor in serie:
            return os.path.join(path_input, wybor)
        print(f"Nieprawidłowa nazwa serii. Wybierz jedną z: {', '.join(serie)}")

def ostrzez_o_nadpisaniu(path_output):
    """Ostrzega o nadpisaniu i pyta o kontynuację, gdy folder nie jest pusty."""
    if os.path.isdir(path_output) and os.listdir(path_output):
        print(
            f"\033[93mOstrzeżenie: folder serii zawiera już dane. "
            f"Istniejące pliki mogą zostać nadpisane: {path_output}\033[0m"
        )
        while True:
            wybor = input("Czy kontynuować? [t/N]: ").strip().casefold()
            if wybor in {"t", "tak"}:
                return True
            if wybor in {"", "n", "nie"}:
                return False
            print("Wpisz 't' (tak) lub 'n' (nie).")
    return True


def utworz_folder_nowej_wersji(path_output):
    folder_wersji = os.path.join(path_output, FOLDER_WERSJI)
    os.makedirs(folder_wersji, exist_ok=True)

    numery_wersji = []
    for nazwa in os.listdir(folder_wersji):
        dopasowanie = re.fullmatch(r"Wersja (\d+)", nazwa)
        if dopasowanie and os.path.isdir(os.path.join(folder_wersji, nazwa)):
            numery_wersji.append(int(dopasowanie.group(1)))

    nastepna_wersja = max(numery_wersji, default=0) + 1
    sciezka_wersji = os.path.join(folder_wersji, f"Wersja {nastepna_wersja}")
    os.makedirs(sciezka_wersji)
    return sciezka_wersji


def przygotuj_materialy(
    konfiguracja,
    pliki_klas,
    zrodla_audio,
    przebiegi_word,
    path_input,
                path_output,
    folder_wersji,
    sciezka_docx,
):
    """Uruchamia etap obrazów, audio i dokumentów po wygenerowaniu przydziałów."""
    liczba_przebiegow = konfiguracja["liczba_przebiegow"]
    path_przebiegi_img = os.path.join(
        folder_wersji, "Rozbicie przebiegów"
    )
    path_przebiegi_audio = os.path.join(
        folder_wersji, "Rozbicie przebiegów audio"
    )

    if not os.path.isdir(path_input):
        raise FileNotFoundError(
            f"Nie znaleziono folderu wybranej serii: {path_input}"
        )
    os.makedirs(path_output, exist_ok=True)

    sukces = wyodrebnij_obrazy_przebiegow(
        sciezka_docx, path_przebiegi_img, przebiegi_word
    )
    przygotuj_pliki_audio(zrodla_audio, path_przebiegi_audio)
    if sukces:
        przetworz_klasy(
            pliki_klas, path_output, path_przebiegi_img, path_przebiegi_audio
        )


def skopiuj_wejscie_i_wyjscie_do_wersji(path_input, path_output, folder_wersji):
    """Zapisuje kopie wejścia i wygenerowanego wyjścia w folderze wersji."""
    shutil.copytree(path_input, os.path.join(folder_wersji, "Input"))

    sciezka_kopii_output = os.path.join(folder_wersji, "Output")

    def pomin_folder_wersji(sciezka, nazwy):
        if os.path.abspath(sciezka) == os.path.abspath(path_output):
            return [FOLDER_WERSJI] if FOLDER_WERSJI in nazwy else []
        return []

    shutil.copytree(
        path_output,
        sciezka_kopii_output,
        ignore=pomin_folder_wersji,
    )


def glowna_funkcja():
    oryginalne_stdout = sys.stdout
    oryginalne_stderr = sys.stderr
    bufor_logu = io.StringIO()
    strumien_stdout = StrumienKonsoliZLogiem(oryginalne_stdout, bufor_logu)
    strumien_stderr = StrumienKonsoliZLogiem(oryginalne_stderr, bufor_logu)
    sys.stdout = strumien_stdout
    sys.stderr = strumien_stderr
    try:
        wykonaj_glowna_funkcje(
            bufor_logu, strumien_stdout, strumien_stderr
        )
    except Exception as error:
        wypisz_blad_na_czerwono(error)
    finally:
        sys.stdout = oryginalne_stdout
        sys.stderr = oryginalne_stderr


def wykonaj_glowna_funkcje(bufor_logu, strumien_stdout, strumien_stderr):
    konfiguracja = wczytaj_konfiguracje(".przebiegi rytmiczne.yml")
    folder_wejsciowy = konfiguracja.get(
        "folder_wejsciowy", FOLDER_PRZEBIEGOW_INPUT
    )
    folder_input = os.path.abspath(
        os.path.join(KATALOG_SKRYPTU, folder_wejsciowy)
    )
    if not os.path.isdir(folder_input):
        raise FileNotFoundError(f"Nie znaleziono folderu wejściowego: {folder_input}")

    path_input = wybierz_serie(folder_input)
    folder_wynikowy = konfiguracja.get("folder_wynikowy", "Wyniki")
    path_output = os.path.abspath(
        os.path.join(
            KATALOG_SKRYPTU,
            folder_wynikowy,
            os.path.basename(os.path.normpath(path_input)),
        )
    )
    if not ostrzez_o_nadpisaniu(path_output):
        print("Przerwano działanie skryptu.")
        return

    sciezka_docx = znajdz_plik_word(path_input)
    przebiegi_word = pobierz_przebiegi_word(sciezka_docx)
    numery_przebiegow = [numer for numer, _ in przebiegi_word]
    konfiguracja["numery_przebiegow"] = numery_przebiegow
    konfiguracja["liczba_przebiegow"] = len(numery_przebiegow)
    print(
        f"Liczba przebiegów odczytana z dokumentu Word: "
        f"{konfiguracja['liczba_przebiegow']}"
    )

    zrodla_audio = znajdz_pliki_audio(
        path_input, numery_przebiegow
    )
    folder_wersji = utworz_folder_nowej_wersji(path_output)
    sciezka_logu = os.path.join(folder_wersji, "Log.txt")
    with open(sciezka_logu, "w", encoding="utf-8") as plik_logu:
        plik_logu.write(bufor_logu.getvalue())
        plik_logu.flush()
        strumien_stdout.podlacz_plik_logu(plik_logu)
        strumien_stderr.podlacz_plik_logu(plik_logu)
        print(f"Log zapisywany w: {sciezka_logu}")
        try:
            wygenerowane_przydzialy = przydziel_numery_dla_klas(
                konfiguracja, path_output, folder_wersji
            )
            przygotuj_materialy(
                konfiguracja,
                wygenerowane_przydzialy,
                zrodla_audio,
                przebiegi_word,
                path_input,
                path_output,
                folder_wersji,
                sciezka_docx,
            )
            skopiuj_wejscie_i_wyjscie_do_wersji(
                path_input, path_output, folder_wersji
            )
            wypisz_raport_koncowy(
                numery_przebiegow,
                wygenerowane_przydzialy,
                os.path.join(path_output, FOLDER_DLA_UCZNIOW),
            )
        except Exception as error:
            wypisz_blad_na_czerwono(error)


if __name__ == "__main__":
    glowna_funkcja()